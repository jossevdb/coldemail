"""
Module voor het exporteren van resultaten naar Excel.
"""
import logging
from datetime import datetime
from pathlib import Path
from typing import List, Dict, Optional
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


class ExcelExporter:
    """Exporteert scraping resultaten naar Excel."""

    def __init__(self, output_dir: str = 'output'):
        """
        Initialiseer de exporter.

        Args:
            output_dir: Directory voor output bestanden
        """
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def export_results(
        self,
        results: List[Dict],
        filename: Optional[str] = None
    ) -> str:
        """
        Exporteer resultaten naar Excel.

        Args:
            results: Lijst van contact dictionaries
            filename: Optionele bestandsnaam (anders timestamp)

        Returns:
            Pad naar het gegenereerde bestand
        """
        if not filename:
            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            filename = f'vlaamse_gemeentes_contacts_{timestamp}.xlsx'

        filepath = self.output_dir / filename

        # Converteer naar DataFrame
        df = pd.DataFrame(results)

        # Sorteer op gemeente en afdeling
        if 'municipality' in df.columns:
            df = df.sort_values(['municipality', 'department'], na_position='last')

        # Voeg datum toe
        df['datum_verzameld'] = datetime.now().strftime('%Y-%m-%d')

        # Herorden kolommen
        column_order = [
            'municipality',
            'province',
            'department',
            'name',
            'email',
            'url',
            'datum_verzameld',
            'context'
        ]

        # Alleen kolommen die bestaan
        columns = [col for col in column_order if col in df.columns]
        df = df[columns]

        # Hernoem kolommen naar Nederlands
        column_names = {
            'municipality': 'Gemeente',
            'province': 'Provincie',
            'department': 'Afdeling',
            'name': 'Naam',
            'email': 'E-mailadres',
            'url': 'Bron URL',
            'datum_verzameld': 'Datum verzameld',
            'context': 'Context'
        }

        df = df.rename(columns=column_names)

        # Schrijf naar Excel
        df.to_excel(filepath, index=False, engine='openpyxl')

        # Pas styling toe
        self._apply_styling(filepath)

        logger.info(f"Resultaten geëxporteerd naar {filepath}")
        logger.info(f"Totaal aantal contacts: {len(df)}")

        # Print statistieken
        self._print_statistics(df)

        return str(filepath)

    def _apply_styling(self, filepath: Path):
        """
        Pas styling toe op Excel bestand.

        Args:
            filepath: Pad naar Excel bestand
        """
        wb = load_workbook(filepath)
        ws = wb.active

        # Header styling
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF')

        for cell in ws[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')

        # Auto-width kolommen
        for column in ws.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)

            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass

            adjusted_width = min(max_length + 2, 50)  # Max 50 karakters breed
            ws.column_dimensions[column_letter].width = adjusted_width

        # Freeze top row
        ws.freeze_panes = 'A2'

        # Sla op
        wb.save(filepath)

    def _print_statistics(self, df: pd.DataFrame):
        """
        Print statistieken over de resultaten.

        Args:
            df: DataFrame met resultaten
        """
        logger.info("\n=== STATISTIEKEN ===")

        if 'Gemeente' in df.columns:
            unique_municipalities = df['Gemeente'].nunique()
            logger.info(f"Unieke gemeentes: {unique_municipalities}")

        if 'Afdeling' in df.columns:
            logger.info("\nContacten per afdeling:")
            dept_counts = df['Afdeling'].value_counts()
            for dept, count in dept_counts.items():
                logger.info(f"  {dept}: {count}")

        if 'Provincie' in df.columns:
            logger.info("\nContacten per provincie:")
            prov_counts = df['Provincie'].value_counts()
            for prov, count in prov_counts.items():
                logger.info(f"  {prov}: {count}")

    def save_checkpoint(
        self,
        results: List[Dict],
        checkpoint_name: str = 'checkpoint'
    ):
        """
        Sla een tussentijdse checkpoint op.

        Args:
            results: Lijst van resultaten tot nu toe
            checkpoint_name: Naam voor checkpoint bestand
        """
        timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
        filename = f'{checkpoint_name}_{timestamp}.xlsx'

        filepath = self.export_results(results, filename)
        logger.info(f"Checkpoint opgeslagen: {filepath}")

        return filepath
